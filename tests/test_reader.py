"""Tests de read_docx_ordered / inspect_docx sur des DOCX FICTIFS créés par python-docx."""

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from dalal_jamm.reports.reader import inspect_docx, read_docx_ordered, read_extras


def save(doc, tmp_path, name="r.docx"):
    path = tmp_path / name
    doc.save(path)
    return path


def test_text_follows_document_order_with_tables_in_place(tmp_path):
    d = Document()
    d.add_paragraph("AVANT")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Nom"
    t.cell(0, 1).text = "Valeur"
    d.add_paragraph("APRES")
    assert read_docx_ordered(save(d, tmp_path)).split("\n") == ["AVANT", "Nom | Valeur", "APRES"]


def test_original_text_is_kept_untouched(tmp_path):
    d = Document()
    d.add_paragraph("Formation isoéchogène : BI-RADS III\u00a0; Surveillance dans 06 mois.")
    assert read_docx_ordered(save(d, tmp_path)) == "Formation isoéchogène : BI-RADS III\u00a0; Surveillance dans 06 mois."


def test_horizontally_merged_cell_is_read_once(tmp_path):
    d = Document()
    t = d.add_table(rows=1, cols=3)
    merged = t.cell(0, 0).merge(t.cell(0, 1))
    merged.text = "FUSION"
    t.cell(0, 2).text = "C"
    assert read_docx_ordered(save(d, tmp_path)) == "FUSION | C"


def test_vertically_merged_cell_is_read_once(tmp_path):
    d = Document()
    t = d.add_table(rows=2, cols=2)
    top = t.cell(0, 0).merge(t.cell(1, 0))
    top.text = "HAUT"
    t.cell(0, 1).text = "a"
    t.cell(1, 1).text = "b"
    assert read_docx_ordered(save(d, tmp_path)).split("\n") == ["HAUT | a", "b"]


def test_nested_table_text_is_not_lost(tmp_path):
    d = Document()
    outer = d.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "interne1"
    inner.cell(0, 1).text = "interne2"
    text = read_docx_ordered(save(d, tmp_path))
    assert "interne1 | interne2" in text
    assert inspect_docx(save(d, tmp_path, "n.docx"))["n_nested_tables"] == 1


def test_content_control_text_is_read(tmp_path):
    d = Document()
    d.add_paragraph("avant")
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>dans le controle</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    )
    d.element.body.insert(len(d.element.body) - 1, sdt)  # avant <w:sectPr>
    assert read_docx_ordered(save(d, tmp_path)).split("\n") == ["avant", "dans le controle"]


def test_empty_document_gives_empty_text(tmp_path):
    assert read_docx_ordered(save(Document(), tmp_path)) == ""


def test_missing_file_raises_without_leaking_the_path(tmp_path):
    with pytest.raises(FileNotFoundError) as exc:
        read_docx_ordered(tmp_path / "NOM_SECRET.docx")
    assert "NOM_SECRET" not in str(exc.value)


def test_corrupted_file_raises(tmp_path):
    bad = tmp_path / "faux.docx"
    bad.write_text("pas un docx " * 10)
    with pytest.raises(Exception):  # noqa: B017 - la classe exacte dépend de la version de python-docx
        read_docx_ordered(bad)


def test_inspect_counts_structure(tmp_path):
    d = Document()
    d.add_paragraph("a")
    d.add_paragraph("b")
    t = d.add_table(rows=1, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1))
    info = inspect_docx(save(d, tmp_path))
    assert info["n_paragraphs"] == 2 and info["n_tables"] == 1 and info["n_merged_cells"] >= 1
    assert info["n_textboxes"] == 0 and info["n_tracked_changes"] == 0 and info["has_header_footer_text"] == 0


def test_textbox_is_flagged_because_its_text_is_not_read(tmp_path):
    d = Document()
    p = d.add_paragraph("corps")
    p._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:txbxContent><w:p><w:r><w:t>ZONE</w:t></w:r></w:p></w:txbxContent></w:r>'
    ))
    path = save(d, tmp_path)
    assert inspect_docx(path)["n_textboxes"] == 1
    assert "ZONE" not in read_docx_ordered(path)  # limite connue, signalée par inspect_docx


def test_tracked_change_is_flagged(tmp_path):
    d = Document()
    p = d.add_paragraph("base ")
    p._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="x" w:date="2026-01-01T00:00:00Z"><w:r><w:t>ajout</w:t></w:r></w:ins>'
    ))
    assert inspect_docx(save(d, tmp_path))["n_tracked_changes"] == 1


def test_header_text_is_flagged_not_mixed_into_body(tmp_path):
    d = Document()
    d.add_paragraph("corps")
    d.sections[0].header.paragraphs[0].text = "En-tête hôpital"
    path = save(d, tmp_path)
    assert inspect_docx(path)["has_header_footer_text"] == 1
    assert read_docx_ordered(path) == "corps"


def textbox_run(text):
    return parse_xml(
        f'<w:r {nsdecls("w")}><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></w:r>'
    )


def test_read_extras_returns_header_footer_and_textbox(tmp_path):
    d = Document()
    p = d.add_paragraph("corps")
    d.sections[0].header.paragraphs[0].text = "En-tête hôpital"
    d.sections[0].footer.paragraphs[0].text = "Pied de page"
    p._p.append(textbox_run("ZONE"))
    assert read_extras(save(d, tmp_path)) == {
        "headers": ["En-tête hôpital"], "footers": ["Pied de page"], "textboxes": ["ZONE"],
    }


def test_read_extras_is_empty_when_there_is_nothing(tmp_path):
    d = Document()
    d.add_paragraph("corps")
    assert read_extras(save(d, tmp_path)) == {"headers": [], "footers": [], "textboxes": []}


def test_header_inherited_by_next_section_is_listed_once(tmp_path):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "H"
    d.add_section()
    assert read_extras(save(d, tmp_path))["headers"] == ["H"]


def test_textbox_written_twice_is_listed_once(tmp_path):
    d = Document()
    p = d.add_paragraph("corps")
    p._p.append(textbox_run("MEME ZONE"))
    p._p.append(textbox_run("MEME ZONE"))
    assert read_extras(save(d, tmp_path))["textboxes"] == ["MEME ZONE"]
